from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from scraper.store import get_conn

OUTPUT_PATH = "output/shirt-sales.xlsx"
MIN_SOLD = 10  # minimum sold listings before surfacing group stats


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _header_style(ws, row: int = 1) -> None:
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws) -> None:
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 50)


def _write_df(ws, df: pd.DataFrame) -> None:
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append(list(row))
    _header_style(ws)
    _autofit(ws)
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
# Data load
# ---------------------------------------------------------------------------

def load_data() -> tuple[pd.DataFrame, pd.DataFrame]:
    conn = get_conn()
    all_df = pd.read_sql_query("SELECT * FROM listings", conn.raw)
    sold_df = pd.read_sql_query("SELECT * FROM listings WHERE sold = 1", conn.raw)
    conn.close()
    return all_df, sold_df


# ---------------------------------------------------------------------------
# Tab builders
# ---------------------------------------------------------------------------

def build_raw(all_df: pd.DataFrame) -> pd.DataFrame:
    cols = [
        "listing_id", "title", "team", "has_player_name", "player_name",
        "shirt_number", "asking_price", "favourites", "view_count",
        "upload_timestamp", "sold", "time_to_sell_days", "size", "condition",
        "season", "is_replica", "is_authentic", "bumped", "bumped_days", "url",
    ]
    return all_df[[c for c in cols if c in all_df.columns]]


def build_by_team(sold_df: pd.DataFrame) -> pd.DataFrame:
    if sold_df.empty:
        return pd.DataFrame(columns=["Team", "Total sold", "Note"])

    grp = sold_df.groupby("team", dropna=True)
    rows = []
    for team, g in grp:
        if len(g) < MIN_SOLD:
            continue
        rows.append({
            "Team": team,
            "Total sold": len(g),
            "Avg price (£)": round(g["asking_price"].mean(), 2),
            "Median price (£)": round(g["asking_price"].median(), 2),
            "Avg time to sell (days)": round(g["time_to_sell_days"].dropna().mean(), 1) if g["time_to_sell_days"].notna().any() else None,
            "Avg favourites": round(g["favourites"].mean(), 1),
            "% with player name": f"{100 * g['has_player_name'].mean():.0f}%",
            "Top player": g["player_name"].dropna().mode().iloc[0] if g["player_name"].notna().any() else None,
        })

    if not rows:
        return pd.DataFrame([{"Note": f"No team has {MIN_SOLD}+ sold listings yet"}])

    return pd.DataFrame(rows).sort_values("Total sold", ascending=False).reset_index(drop=True)


def build_named_vs_blank(sold_df: pd.DataFrame) -> pd.DataFrame:
    if sold_df.empty:
        return pd.DataFrame(columns=["Type", "Note"])

    def stats(df: pd.DataFrame, label: str) -> dict:
        return {
            "Type": label,
            "Count sold": len(df),
            "Avg price (£)": round(df["asking_price"].mean(), 2),
            "Median price (£)": round(df["asking_price"].median(), 2),
            "Avg time to sell (days)": round(df["time_to_sell_days"].dropna().mean(), 1) if df["time_to_sell_days"].notna().any() else None,
            "Avg favourites": round(df["favourites"].mean(), 1),
        }

    overall = [
        stats(sold_df[sold_df["has_player_name"] == 1], "Named — all teams"),
        stats(sold_df[sold_df["has_player_name"] == 0], "Blank — all teams"),
    ]

    per_team = []
    for team in sorted(sold_df["team"].dropna().unique()):
        t = sold_df[sold_df["team"] == team]
        named = t[t["has_player_name"] == 1]
        blank = t[t["has_player_name"] == 0]
        if len(named) >= MIN_SOLD:
            per_team.append(stats(named, f"{team} — named"))
        if len(blank) >= MIN_SOLD:
            per_team.append(stats(blank, f"{team} — blank"))

    return pd.DataFrame(overall + per_team)


def build_player_names(sold_df: pd.DataFrame) -> pd.DataFrame:
    named = sold_df[sold_df["player_name"].notna()]
    if named.empty:
        return pd.DataFrame(columns=["Player", "Note"])

    grp = named.groupby("player_name")
    rows = []
    for player, g in grp:
        if len(g) < 3:
            continue
        rows.append({
            "Player": player,
            "Times sold": len(g),
            "Teams": ", ".join(sorted(g["team"].dropna().unique())),
            "Avg price (£)": round(g["asking_price"].mean(), 2),
            "Avg favourites": round(g["favourites"].mean(), 1),
            "Avg time to sell (days)": round(g["time_to_sell_days"].dropna().mean(), 1) if g["time_to_sell_days"].notna().any() else None,
        })

    if not rows:
        return pd.DataFrame([{"Note": "No player has 3+ sold listings yet"}])

    return pd.DataFrame(rows).sort_values("Times sold", ascending=False).reset_index(drop=True)


def build_price_bands(sold_df: pd.DataFrame) -> pd.DataFrame:
    if sold_df.empty:
        return pd.DataFrame(columns=["Price band", "Note"])

    bands = [(0, 10), (10, 20), (20, 30), (30, 50), (50, 999)]
    rows = []
    for lo, hi in bands:
        label = f"£{lo}–{hi}" if hi < 999 else f"£{lo}+"
        g = sold_df[(sold_df["asking_price"] >= lo) & (sold_df["asking_price"] < hi)]
        rows.append({
            "Price band": label,
            "Sold count": len(g),
            "Avg time to sell (days)": round(g["time_to_sell_days"].dropna().mean(), 1) if g["time_to_sell_days"].notna().any() else None,
            "Avg favourites": round(g["favourites"].mean(), 1) if not g.empty else None,
            "% with player name": f"{100 * g['has_player_name'].mean():.0f}%" if not g.empty else None,
        })

    return pd.DataFrame(rows)


def build_interesting_stats(all_df: pd.DataFrame, sold_df: pd.DataFrame) -> pd.DataFrame:
    rows = []

    def add(stat, value):
        rows.append({"Stat": stat, "Value": value})

    add("Total listings seen", len(all_df))
    add("Total sold", len(sold_df))
    add("Total active", int((all_df["sold"] == 0).sum()))

    if not sold_df.empty:
        fastest = sold_df.loc[sold_df["time_to_sell_days"].idxmin()] if sold_df["time_to_sell_days"].notna().any() else None
        if fastest is not None:
            add("Fastest sold shirt", f"{fastest['title']} ({fastest['time_to_sell_days']:.1f} days)")

        priciest = sold_df.loc[sold_df["asking_price"].idxmax()]
        add("Highest priced sold", f"£{priciest['asking_price']} — {priciest['title']}")

    if not all_df.empty:
        most_faved = all_df.loc[all_df["favourites"].idxmax()]
        add("Most favourited listing", f"{most_faved['favourites']} faves — {most_faved['title']}")

    if not sold_df.empty:
        top_team = sold_df.groupby("team").size().idxmax() if sold_df["team"].notna().any() else None
        if top_team:
            add("Team with most sold", top_team)

        top_player = sold_df["player_name"].dropna().mode()
        if not top_player.empty:
            add("Most common player name sold", top_player.iloc[0])

        named = sold_df[sold_df["has_player_name"] == 1]["time_to_sell_days"].dropna().mean()
        blank = sold_df[sold_df["has_player_name"] == 0]["time_to_sell_days"].dropna().mean()
        if named and blank and blank > 0:
            ratio = named / blank
            add("Named vs blank sell speed ratio", f"{ratio:.2f}x (named {'faster' if ratio < 1 else 'slower'})")

    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def export(output_path: str = OUTPUT_PATH) -> None:
    print("Loading data from Turso...")
    all_df, sold_df = load_data()
    print(f"  {len(all_df)} total listings, {len(sold_df)} sold")

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        tabs = [
            ("Raw Data",        build_raw(all_df)),
            ("By Team",         build_by_team(sold_df)),
            ("Named vs Blank",  build_named_vs_blank(sold_df)),
            ("Player Names",    build_player_names(sold_df)),
            ("Price Bands",     build_price_bands(sold_df)),
            ("Interesting Stats", build_interesting_stats(all_df, sold_df)),
        ]
        for sheet_name, df in tabs:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            ws = writer.sheets[sheet_name]
            _header_style(ws)
            _autofit(ws)
            ws.freeze_panes = "A2"

    print(f"Spreadsheet written to {output_path}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else OUTPUT_PATH
    export(output_path=out)
